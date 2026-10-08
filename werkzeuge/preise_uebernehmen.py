#!/usr/bin/env python3
"""Übernimmt die interne Bezugsquellen-Recherche in Verkaufspreise und eine interne Einkaufsliste.

Liest  /mnt/project-files/katalog/intern/eingabe-*.csv und bezug-*.csv (nicht im Repo, enthält Händler und EK).
Schreibt daten/preise-vergleich.csv (nur Artikelnummer und Verkaufspreis, öffentlich unbedenklich)
und     /mnt/project-files/katalog/intern/bezugsquellen.xlsx (intern, für Bestellungen).

Aufruf: python3 werkzeuge/preise_uebernehmen.py [aufschlag_prozent]   (Standard 10)
"""
import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
INTERN = Path("/mnt/project-files/katalog/intern")
AUFSCHLAG = float(sys.argv[1]) if len(sys.argv) > 1 else 10.0
WENIG_MARGE = 2.00  # € Aufschlag pro Stück, darunter wird es in der Liste markiert
# Preise, die nicht als Stückpreis dieses Artikels gelten dürfen: bleiben intern, Website zeigt „auf Anfrage“.
NICHT_AUF_WEBSITE = {
    "BF-V-5A07C449": "Preis gilt für M16, Artikel ist M24 (71,02 € laut Notiz)",
}


def sperrgrund(sku, b):
    if sku in NICHT_AUF_WEBSITE:
        return NICHT_AUF_WEBSITE[sku]
    bezug = b.get("preis_bezug", "").lower()
    if "packung" in bezug or "paar" in bezug:
        return "Packungspreis, kein Stückpreis"
    return ""


def zahl(text):
    try:
        return float(str(text).replace(",", "."))
    except ValueError:
        return None


def main():
    eingabe = {}
    for f in sorted(INTERN.glob("eingabe-*.csv")):
        for r in csv.DictReader(f.open(encoding="utf-8")):
            eingabe[r["sku"]] = r
    bezug = {}
    for f in sorted(INTERN.glob("bezug-*.csv")):
        for r in csv.DictReader(f.open(encoding="utf-8")):
            bezug[r["sku"]] = r

    zeilen, vk = [], {}
    for sku, e in eingabe.items():
        b = bezug.get(sku, {})
        ek = zahl(b.get("preis_netto_eur", ""))
        preis = round(ek * (1 + AUFSCHLAG / 100), 2) if ek else None
        grund = sperrgrund(sku, b) if preis else ""
        if preis and not grund:
            vk[sku] = preis
        zeilen.append([sku, e["hersteller"], e["serie_artikel"], b.get("bezugsquelle", ""), b.get("bestell_url", ""),
                       ek, preis, round(preis - ek, 2) if preis else None, b.get("preis_bezug", ""),
                       b.get("preis_quelle_url", ""), "ja" if sku in vk else ("nein: " + grund if grund else "nein: kein Preis"),
                       b.get("bild_verfuegbar", ""), b.get("cad_verfuegbar", ""), b.get("notiz", "")])

    with (ROOT / "daten/preise-vergleich.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["sku", "preis_netto"])
        w.writerows([s, f"{p:.2f}"] for s, p in sorted(vk.items()))

    wb = Workbook()
    ws = wb.active
    ws.title = "Bezugsquellen"
    kopf = ["Art.-Nr. Shop", "Hersteller", "Serie / Artikel", "Bezugsquelle", "Bestell-Link", "EK netto €",
            f"VK netto € (+{AUFSCHLAG:g} %)", "Aufschlag €", "Preis bezieht sich auf", "Preisquelle", "Preis auf Website", "Fotos beim Hersteller",
            "CAD beim Hersteller", "Notiz"]
    ws.append(kopf)
    for z in sorted(zeilen, key=lambda z: (z[1], z[2])):
        ws.append(z)
    breiten = [16, 20, 34, 24, 40, 11, 14, 11, 26, 40, 22, 12, 12, 50]
    rot = PatternFill("solid", fgColor="FDE2E1")
    for c, b in enumerate(breiten, 1):
        ws.column_dimensions[get_column_letter(c)].width = b
        ws.cell(1, c).font = Font(bold=True, color="FFFFFF")
        ws.cell(1, c).fill = PatternFill("solid", fgColor="1B2836")
    for row in ws.iter_rows(min_row=2):
        for idx in (4, 9):
            if row[idx].value and str(row[idx].value).startswith("http"):
                row[idx].hyperlink = row[idx].value
        if row[7].value is not None and row[7].value < WENIG_MARGE:
            row[7].fill = rot
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(kopf))}{len(zeilen) + 1}"
    wb.save(INTERN / "bezugsquellen.xlsx")

    mit_quelle = sum(1 for z in zeilen if z[3])
    knapp = sum(1 for z in zeilen if z[0] in vk and z[7] < WENIG_MARGE)
    print(f"{len(zeilen)} Produkte, {mit_quelle} mit Bezugsquelle, {len(vk)} mit Preis, {knapp} mit unter {WENIG_MARGE:.2f} € Aufschlag")


if __name__ == "__main__":
    main()
