#!/usr/bin/env python3
"""Führt die Herstellerrecherche (katalog/recherche/teil-*.csv) zu einer Excel-Auswahlliste zusammen.

Aufruf: python3 werkzeuge/build_auswahlliste.py <recherche-ordner> <ziel.xlsx>
"""
import csv
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SPALTEN = [
    ("Nr", 6), ("Aufnehmen?", 12), ("Produktwelt", 22), ("Hersteller", 24), ("Serie / Artikel", 30),
    ("Kurzbeschreibung", 48), ("Material", 32), ("Gewinde / Anschluss", 20), ("Maß / Ø (mm)", 16),
    ("Traglast (N)", 14), ("Weitere Daten", 26), ("Bodenbefestigung", 16), ("Zertifizierung", 24),
    ("Varianten", 10), ("Quelle", 50),
]

# Feldnamen der einzelnen Recherche-Dateien → gemeinsame Spalten.
ZUORDNUNG = {
    "Maß / Ø (mm)": ["rohrmass_oder_durchmesser_mm", "teller_durchmesser_mm", "durchmesser_mm"],
    "Weitere Daten": [("Verstellbereich", "verstellbereich_mm"), ("Gelenkwinkel °", "gelenkwinkel_grad"), ("Dämpfung", "daempfung_haerte")],
}

NAVY = "1B2836"
GRUEN = "4FA82E"


def zeile(r):
    mass = next((r[k] for k in ZUORDNUNG["Maß / Ø (mm)"] if r.get(k)), "")
    weitere = "; ".join(f"{label}: {r[k]}" for label, k in ZUORDNUNG["Weitere Daten"] if r.get(k))
    return [
        "", r.get("produktwelt", ""), r.get("hersteller", ""), r.get("serie_artikel", ""), r.get("kurzbeschreibung", ""),
        r.get("material", ""), r.get("anschluss_gewinde", ""), mass, r.get("traglast_n", ""), weitere,
        r.get("bodenbefestigung", ""), r.get("zertifizierung", ""), r.get("anzahl_varianten", ""), r.get("url", ""),
    ]


def main(ordner: Path, ziel: Path):
    daten = []
    for f in sorted(ordner.glob("teil-*.csv")):
        daten += [zeile(r) for r in csv.DictReader(f.open(encoding="utf-8"))]
    daten.sort(key=lambda z: (z[1], z[2].lower(), z[3]))

    wb = Workbook()
    ws = wb.active
    ws.title = "Auswahlliste"
    ws.append([s for s, _ in SPALTEN])
    for i, z in enumerate(daten, 1):
        ws.append([i] + z)
    for c, (_, breite) in enumerate(SPALTEN, 1):
        ws.column_dimensions[get_column_letter(c)].width = breite
        kopf = ws.cell(1, c)
        kopf.font = Font(bold=True, color="FFFFFF")
        kopf.fill = PatternFill("solid", fgColor=NAVY)
        kopf.alignment = Alignment(vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in (6, 7, 11, 13))
        url = row[-1]
        if url.value and str(url.value).startswith("http"):
            url.hyperlink = str(url.value)
            url.font = Font(color="3D8A22", underline="single")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(SPALTEN))}{len(daten) + 1}"
    dv = DataValidation(type="list", formula1='"ja,nein,prüfen"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"B2:B{len(daten) + 1}")

    ue = wb.create_sheet("Übersicht", 0)
    ue["A1"] = "Vergleichsprodukte für beinundfuß.de"
    ue["A1"].font = Font(bold=True, size=14, color=NAVY)
    ue["A2"] = (f"{len(daten)} Einträge (je Serie und Größe), recherchiert am 08.10.2026 auf Hersteller- und Händlerseiten. "
                "Interne Auswahlliste: Texte und Bilder nicht übernehmen, Daten vor Aufnahme beim Hersteller anfragen.")
    ue["A2"].alignment = Alignment(wrap_text=True)
    ue.merge_cells("A2:D2")
    ue.row_dimensions[2].height = 45
    ue["A3"] = "So geht's: Im Blatt „Auswahlliste“ in Spalte „Aufnehmen?“ ja / nein / prüfen wählen und die Datei zurückschicken."
    r = 5
    for titel, zaehler in (("Nach Produktwelt", Counter(z[1] for z in daten)),
                           ("Nach Hersteller", Counter(z[2].split(" (")[0] for z in daten))):
        ue.cell(r, 1, titel).font = Font(bold=True, color="FFFFFF")
        ue.cell(r, 1).fill = PatternFill("solid", fgColor=GRUEN)
        ue.cell(r, 2, "Einträge").font = Font(bold=True, color="FFFFFF")
        ue.cell(r, 2).fill = PatternFill("solid", fgColor=GRUEN)
        r += 1
        for name, n in zaehler.most_common():
            ue.cell(r, 1, name)
            ue.cell(r, 2, n)
            r += 1
        r += 1
    ue.column_dimensions["A"].width = 44
    ue.column_dimensions["B"].width = 12
    ue.column_dimensions["C"].width = 20
    ue.column_dimensions["D"].width = 30

    wb.save(ziel)
    print(f"{ziel}: {len(daten)} Einträge")


if __name__ == "__main__":
    main(Path(sys.argv[1]), Path(sys.argv[2]))
